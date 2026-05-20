# coding=utf-8
"""
@project: MaxKB
@file： test_template_parser.py
@desc: Unit tests for finance.service.template_parser.

Pure-Python tests — no Django db needed. We build the test .docx on the
fly with python-docx (already a project dep) and feed its bytes to the
parser.
"""

from io import BytesIO
from unittest import TestCase, skipUnless

from finance.service.template_parser import _infer_type, extract_placeholders


try:  # pragma: no cover — env check
    import docx as _python_docx  # noqa: F401
    import docxtpl as _docxtpl  # noqa: F401

    _HAS_DOCX_STACK = True
except Exception:  # noqa: BLE001
    _HAS_DOCX_STACK = False


def _build_docx_bytes(text: str) -> bytes:
    """Build a one-paragraph docx containing `text` and return its bytes."""
    import docx  # local: only when the stack is available

    document = docx.Document()
    document.add_paragraph(text)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


class InferTypeTest(TestCase):
    """Heuristic typing should be predictable."""

    def test_date_suffix(self):
        self.assertEqual(_infer_type("issue_date"), "date")
        self.assertEqual(_infer_type("signed_at"), "date")
        self.assertEqual(_infer_type("expiry_deadline"), "date")

    def test_number_suffix(self):
        self.assertEqual(_infer_type("loan_amount"), "number")
        self.assertEqual(_infer_type("share_count"), "number")
        self.assertEqual(_infer_type("total_quantity"), "number")

    def test_long_text_suffix(self):
        self.assertEqual(_infer_type("risk_analysis"), "long_text")
        self.assertEqual(_infer_type("project_summary"), "long_text")
        self.assertEqual(_infer_type("management_remark"), "long_text")

    def test_default_text(self):
        self.assertEqual(_infer_type("company_name"), "text")
        self.assertEqual(_infer_type("industry"), "text")


@skipUnless(_HAS_DOCX_STACK, "python-docx + docxtpl not installed")
class ExtractPlaceholdersTest(TestCase):
    def test_returns_metadata_for_each_variable(self):
        body = "公司名称: {{ company_name }}，金额: {{ loan_amount }}，签约日期: {{ signed_at }}."
        placeholders = extract_placeholders(_build_docx_bytes(body))
        by_key = {p["key"]: p for p in placeholders}
        self.assertIn("company_name", by_key)
        self.assertIn("loan_amount", by_key)
        self.assertIn("signed_at", by_key)

        self.assertEqual(by_key["company_name"]["type"], "text")
        self.assertEqual(by_key["loan_amount"]["type"], "number")
        self.assertEqual(by_key["signed_at"]["type"], "date")

        self.assertEqual(by_key["company_name"]["label"], "公司名称")
        self.assertEqual(by_key["loan_amount"]["label"], "金额")
        self.assertEqual(by_key["signed_at"]["label"], "签约日期")

        for p in placeholders:
            # Each placeholder carries the full metadata schema.
            self.assertEqual(set(p.keys()), {"key", "label", "type", "required", "ai_hint", "enum_options"})
            self.assertFalse(p["required"])
            self.assertEqual(p["ai_hint"], f"请填写{p['label']}，保持与模板上下文一致。")
            self.assertEqual(p["enum_options"], [])

    def test_infers_label_from_text_before_placeholder_and_defaults_optional(self):
        body = "借款人名称：{{ borrower_name }}，贷款金额：{{ loan_amount }}。"
        placeholders = extract_placeholders(_build_docx_bytes(body))
        by_key = {p["key"]: p for p in placeholders}

        self.assertEqual(by_key["borrower_name"]["label"], "借款人名称")
        self.assertEqual(by_key["loan_amount"]["label"], "贷款金额")
        self.assertFalse(by_key["borrower_name"]["required"])
        self.assertFalse(by_key["loan_amount"]["required"])
        self.assertEqual(by_key["borrower_name"]["ai_hint"], "请填写借款人名称，保持与模板上下文一致。")
        self.assertEqual(by_key["loan_amount"]["ai_hint"], "请填写贷款金额，保持与模板上下文一致。")

    def test_deduplicates_repeats(self):
        body = "{{ company_name }} - {{ company_name }} - {{ company_name }}"
        placeholders = extract_placeholders(_build_docx_bytes(body))
        keys = [p["key"] for p in placeholders]
        self.assertEqual(keys, ["company_name"])

    def test_empty_template(self):
        placeholders = extract_placeholders(_build_docx_bytes("No placeholders here."))
        self.assertEqual(placeholders, [])

    def test_long_text_heuristic(self):
        body = "风险分析：{{ risk_analysis }}"
        placeholders = extract_placeholders(_build_docx_bytes(body))
        self.assertEqual(len(placeholders), 1)
        self.assertEqual(placeholders[0]["type"], "long_text")
