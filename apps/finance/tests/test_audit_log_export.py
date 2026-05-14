# coding=utf-8
"""
    @project: MaxKB
    @file:   test_audit_log_export.py
    @desc:   Unit tests for FinanceAuditLogExportView helpers (Gate 6 Track A4).

    Like test_audit_log_view.py these tests use SimpleTestCase and avoid
    the database. The behaviours covered:
      - filename generation: workspace id is sanitised and timestamp
        present
      - CSV row iterator: yields BOM then header then data rows, payload
        is JSON-encoded with non-ASCII preserved
      - export cap constant is exposed and sensible (50K)
      - route is wired into the finance URL conf
"""
from types import SimpleNamespace

from django.test import SimpleTestCase

from finance.views.audit_log import (
    _CSV_HEADERS,
    _EXPORT_MAX_ROWS,
    _XLSX_CONTENT_TYPE,
    _build_xlsx_bytes,
    _filename_for_export,
    _iter_csv_rows,
)


class FilenameForExportTest(SimpleTestCase):
    def test_filename_includes_workspace_and_extension(self):
        name = _filename_for_export('default')
        self.assertTrue(name.startswith('finance_audit_default_'))
        self.assertTrue(name.endswith('.csv'))

    def test_unsafe_workspace_chars_stripped(self):
        # Spaces, slashes, etc. must not land in the filename header.
        name = _filename_for_export('weird ws/with..bad?chars')
        # Hyphens / underscores survive; everything else is stripped.
        self.assertNotIn(' ', name)
        self.assertNotIn('/', name)
        self.assertNotIn('?', name)

    def test_falls_back_to_default_for_empty_id(self):
        name = _filename_for_export('')
        self.assertIn('finance_audit_default_', name)


class IterCsvRowsTest(SimpleTestCase):
    def _fake_row(self, **overrides):
        defaults = dict(
            created_at=SimpleNamespace(isoformat=lambda: '2026-05-13T08:00:00+00:00'),
            actor_id='11111111-1111-1111-1111-111111111111',
            action='UPDATE',
            target_type='PROJECT',
            target_id='22222222-2222-2222-2222-222222222222',
            ip='127.0.0.1',
            user_agent='UA/1.0',
            payload={'path': '/api/finance/...', 'note': '中文测试'},
        )
        defaults.update(overrides)
        return SimpleNamespace(**defaults)

    class _FakeQS:
        """Mimics a Django queryset for ``iterator()``."""
        def __init__(self, rows):
            self._rows = rows

        def iterator(self, chunk_size=500):
            yield from self._rows

    def test_iter_yields_bom_header_and_rows(self):
        rows = [self._fake_row(), self._fake_row(target_type='SMTP_CONFIG')]
        chunks = list(_iter_csv_rows(self._FakeQS(rows)))
        # Concatenate to one bytes blob to inspect the document shape.
        blob = b''.join(chunks).decode('utf-8-sig')  # strips BOM
        lines = [line for line in blob.splitlines() if line.strip()]
        # 1 header + 2 data lines
        self.assertEqual(len(lines), 3)
        self.assertIn('created_at', lines[0])
        self.assertIn('actor_id', lines[0])
        self.assertIn('UPDATE', lines[1])
        self.assertIn('SMTP_CONFIG', lines[2])

    def test_iter_preserves_non_ascii_payload(self):
        rows = [self._fake_row()]
        chunks = list(_iter_csv_rows(self._FakeQS(rows)))
        blob = b''.join(chunks).decode('utf-8-sig')
        # Chinese payload must survive — ensure_ascii=False on json.dumps.
        self.assertIn('中文测试', blob)

    def test_iter_handles_empty_payload(self):
        rows = [self._fake_row(payload=None)]
        chunks = list(_iter_csv_rows(self._FakeQS(rows)))
        blob = b''.join(chunks).decode('utf-8-sig')
        # Payload column should be the JSON empty-object string.
        # We don't pin the exact column index, but '{}' must appear.
        self.assertIn('{}', blob)

    def test_iter_handles_unjsonable_payload(self):
        # Force a TypeError in json.dumps via a non-serialisable object.
        class _Bomb:
            def __repr__(self):  # default=str will convert via repr
                return '<_Bomb>'

        # default=str saves us — but if a payload was somehow circular
        # (CIRCULAR_REF detection raises ValueError) we still expect an
        # empty string, not a crash.
        rows = [self._fake_row(payload={'b': _Bomb()})]
        chunks = list(_iter_csv_rows(self._FakeQS(rows)))
        # Did not raise; concat works.
        b''.join(chunks)


class ConstantsTest(SimpleTestCase):
    def test_export_cap_is_50k(self):
        self.assertEqual(_EXPORT_MAX_ROWS, 50_000)

    def test_csv_headers_match_expected_columns(self):
        # Compliance reviewers eyeball this — the column count + order
        # is part of the contract.
        self.assertEqual(len(_CSV_HEADERS), 8)
        self.assertTrue(any('created_at' in h for h in _CSV_HEADERS))
        self.assertTrue(any('Payload' in h for h in _CSV_HEADERS))


class AuditLogExportRouteTest(SimpleTestCase):
    def test_export_url_resolves(self):
        from django.urls import reverse

        url = reverse(
            'finance:audit_log_export',
            kwargs={'workspace_id': 'default'},
        )
        self.assertEqual(url, '/api/finance/workspace/default/audit-log/export')


# ---- Gate 7 Track A2: xlsx export tests ----


class XlsxFilenameTest(SimpleTestCase):
    def test_xlsx_filename_uses_xlsx_extension(self):
        name = _filename_for_export('default', ext='xlsx')
        self.assertTrue(name.startswith('finance_audit_default_'))
        self.assertTrue(name.endswith('.xlsx'))

    def test_csv_default_unchanged(self):
        # Existing callers passing no ext must still get csv.
        name = _filename_for_export('default')
        self.assertTrue(name.endswith('.csv'))

    def test_unsupported_ext_falls_back_to_csv(self):
        name = _filename_for_export('default', ext='pdf')
        self.assertTrue(name.endswith('.csv'))


class XlsxBuildTest(SimpleTestCase):
    """
    Round-trip an in-memory workbook to verify shape: openpyxl can read
    its own output and we get the rows we wrote.
    """

    def _fake_row(self, **overrides):
        defaults = dict(
            created_at=SimpleNamespace(isoformat=lambda: '2026-05-13T08:00:00+00:00'),
            actor_id='11111111-1111-1111-1111-111111111111',
            action='UPDATE',
            target_type='PROJECT',
            target_id='22222222-2222-2222-2222-222222222222',
            ip='127.0.0.1',
            user_agent='UA/1.0',
            payload={'path': '/api/finance/...', 'note': '中文测试'},
        )
        defaults.update(overrides)
        return SimpleNamespace(**defaults)

    class _FakeQS:
        def __init__(self, rows):
            self._rows = rows

        def iterator(self, chunk_size=500):
            yield from self._rows

    def test_xlsx_bytes_are_valid_workbook(self):
        rows = [self._fake_row(), self._fake_row(target_type='SMTP_CONFIG')]
        blob = _build_xlsx_bytes(self._FakeQS(rows))
        # xlsx files are zip archives starting with PK\x03\x04
        self.assertTrue(blob[:2] == b'PK', 'output must be a zip-shaped xlsx')

        # Re-open and inspect.
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(blob), read_only=True)
        ws = wb['finance_audit']
        rows_iter = list(ws.iter_rows(values_only=True))
        # 1 header + 2 data
        self.assertEqual(len(rows_iter), 3)
        # Header label sanity
        self.assertIn('created_at', str(rows_iter[0][0]))
        # Action column on row 1
        self.assertEqual(rows_iter[1][2], 'UPDATE')
        # target_type column on row 2
        self.assertEqual(rows_iter[2][3], 'SMTP_CONFIG')

    def test_xlsx_payload_column_preserves_chinese(self):
        rows = [self._fake_row()]
        blob = _build_xlsx_bytes(self._FakeQS(rows))
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(blob), read_only=True)
        ws = wb['finance_audit']
        data_rows = list(ws.iter_rows(values_only=True))
        payload_cell = data_rows[1][-1]  # last column is payload
        self.assertIn('中文测试', payload_cell)

    def test_xlsx_content_type_constant(self):
        # Sanity: the type string is the canonical Office XML one.
        self.assertEqual(
            _XLSX_CONTENT_TYPE,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
