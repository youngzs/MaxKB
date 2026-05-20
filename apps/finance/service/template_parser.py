# coding=utf-8
"""
@project: MaxKB
@file： template_parser.py
@desc: docx → placeholder-metadata extraction.

Uses docxtpl's DocxTemplate.get_undeclared_template_variables() to find
every Jinja-style variable referenced in the document body, headers and
footers. Returns a list of placeholder dicts shaped exactly like
DocumentTemplate.placeholders:

    [{key, label, type, required, ai_hint, enum_options}, ...]

`type` is inferred from the variable name with a small set of heuristics
(see _infer_type) — the user can override it later from the metadata
editor.

Pure function; no Django or OSS dependency, so unit tests can exercise it
by feeding raw docx bytes built with python-docx.
"""

from __future__ import annotations

from io import BytesIO
import re
from typing import Iterable


_NUMBER_SUFFIXES = ("_amount", "_count", "_number", "_qty", "_quantity", "_total")
_DATE_SUFFIXES = ("_date", "_at", "_time", "_deadline", "_expiry")
_LONG_TEXT_SUFFIXES = (
    "_desc",
    "_description",
    "_summary",
    "_analysis",
    "_reason",
    "_notes",
    "_remark",
    "_content",
)


def _infer_type(key: str) -> str:
    """Heuristic: variable suffix → placeholder type."""
    lowered = key.lower()
    if lowered.endswith(_DATE_SUFFIXES):
        return "date"
    if lowered.endswith(_NUMBER_SUFFIXES):
        return "number"
    if lowered.endswith(_LONG_TEXT_SUFFIXES):
        return "long_text"
    return "text"


_KEY_WORDS_ZH = {
    "abs": "资产证券化",
    "address": "地址",
    "amount": "金额",
    "analysis": "分析",
    "annual": "年度",
    "asset": "资产",
    "bank": "银行",
    "borrower": "借款人",
    "brief": "摘要",
    "business": "业务",
    "company": "公司",
    "contact": "联系人",
    "contract": "合同",
    "count": "数量",
    "credit": "授信",
    "currency": "币种",
    "date": "日期",
    "deadline": "截止日期",
    "debt": "债务",
    "desc": "描述",
    "description": "描述",
    "equity": "股权",
    "expiry": "到期日",
    "finance": "融资",
    "financing": "融资",
    "guarantee": "担保",
    "industry": "行业",
    "interest": "利息",
    "issue": "出具",
    "issuer": "发行人",
    "lender": "贷款人",
    "loan": "贷款",
    "manager": "经理",
    "meeting": "会议",
    "name": "名称",
    "notes": "备注",
    "number": "编号",
    "party": "方",
    "period": "期限",
    "project": "项目",
    "purpose": "用途",
    "qty": "数量",
    "quantity": "数量",
    "rate": "利率",
    "reason": "原因",
    "region": "地区",
    "remark": "备注",
    "repayment": "还款",
    "report": "报告",
    "risk": "风险",
    "signed": "签署",
    "summary": "摘要",
    "term": "期限",
    "time": "时间",
    "title": "标题",
    "total": "总计",
    "type": "类型",
}

_PLACEHOLDER_RE = re.compile(r"{{\s*([A-Za-z_][A-Za-z0-9_]*)\s*(?:\|[^}]*)?}}")
_LABEL_SPLIT_RE = re.compile(r"[\n\r\t,，;；。.!！?？]")
_TRIM_LABEL_CHARS = " \t\r\n:：-—_、，,;；。.!！?？()（）[]【】<>《》\"“”'"


def infer_label_from_key(key: str) -> str:
    """Best-effort Chinese label from a snake_case/camelCase placeholder key."""
    normalized = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", key).lower()
    words = [w for w in re.split(r"[_\-\s]+", normalized) if w]
    translated = [_KEY_WORDS_ZH.get(w, w) for w in words]
    label = "".join(translated).strip()
    return label or key


def build_ai_hint(label: str) -> str:
    """Default one-sentence hint shown to AI fill; intentionally conservative."""
    return f"请填写{label}，保持与模板上下文一致。"


def _clean_label(value: str) -> str:
    cleaned = re.sub(r"{{.*?}}", "", value or "").strip(_TRIM_LABEL_CHARS)
    cleaned = re.sub(r"\s+", "", cleaned)
    if len(cleaned) > 40:
        cleaned = cleaned[-40:].strip(_TRIM_LABEL_CHARS)
    return cleaned


def _label_from_text_before_placeholder(text: str, match_start: int) -> str:
    prefix = text[:match_start].rstrip(_TRIM_LABEL_CHARS)
    if not prefix:
        return ""
    candidate = _LABEL_SPLIT_RE.split(prefix)[-1]
    return _clean_label(candidate)


def _labels_from_plain_text(text: str, wanted_keys: set[str]) -> dict[str, str]:
    labels: dict[str, str] = {}
    for match in _PLACEHOLDER_RE.finditer(text or ""):
        key = match.group(1)
        if key not in wanted_keys or key in labels:
            continue
        label = _label_from_text_before_placeholder(text, match.start())
        if label:
            labels[key] = label
    return labels


def _iter_docx_paragraph_text(document) -> Iterable[str]:
    for paragraph in document.paragraphs:
        if paragraph.text:
            yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    yield text
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                if paragraph.text:
                    yield paragraph.text
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            yield text


def _extract_labels_from_docx(docx_bytes: bytes, wanted_keys: set[str]) -> dict[str, str]:
    try:
        import docx

        document = docx.Document(BytesIO(docx_bytes))
    except Exception:  # noqa: BLE001
        return {}

    labels: dict[str, str] = {}
    for text in _iter_docx_paragraph_text(document):
        for key, label in _labels_from_plain_text(text, wanted_keys).items():
            labels.setdefault(key, label)
    return labels


def extract_plain_text(docx_bytes: bytes, *, limit: int = 6000) -> str:
    """Extract readable text from a docx for AI metadata prompts."""
    try:
        import docx

        document = docx.Document(BytesIO(docx_bytes))
    except Exception:  # noqa: BLE001
        return ""
    parts = [text.strip() for text in _iter_docx_paragraph_text(document) if text.strip()]
    return "\n".join(parts)[:limit]


def _build_placeholder(key: str, label: str | None = None) -> dict:
    display_label = _clean_label(label or "") or infer_label_from_key(key)
    return {
        "key": key,
        "label": display_label,
        "type": _infer_type(key),
        "required": False,
        "ai_hint": build_ai_hint(display_label),
        "enum_options": [],
    }


def extract_placeholders(docx_bytes: bytes) -> list[dict]:
    """
    Scan a docx file for Jinja-style placeholders {{ var }} / {% block %}
    using docxtpl.

    Args:
        docx_bytes: full bytes of a .docx file.

    Returns:
        Stable-ordered list of placeholder dicts (one per unique variable).
    """
    # Local import: docxtpl pulls in python-docx + lxml which we don't want
    # to load on every Django boot.
    from docxtpl import DocxTemplate

    template = DocxTemplate(BytesIO(docx_bytes))
    # docxtpl returns a Python set; sort for deterministic output.
    variables: Iterable[str] = template.get_undeclared_template_variables() or set()
    seen: list[str] = sorted({str(v) for v in variables if v})
    labels_by_key = _extract_labels_from_docx(docx_bytes, set(seen))
    return [_build_placeholder(key, labels_by_key.get(key)) for key in seen]
