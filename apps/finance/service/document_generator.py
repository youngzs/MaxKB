# coding=utf-8
"""
    @project: MaxKB
    @file： document_generator.py
    @desc: Render a DocumentTemplate against a placeholder-value dict, store
    the resulting docx in OSS, and (separately) convert it to HTML for
    preview via `mammoth`.

    `trigger_generation` is intentionally synchronous in Gate 3 — it runs in
    the same process as the POST request. Gate 5 will wrap it in a Celery
    task; the function signature already takes only an id so it can become
    `@shared_task` later without callsite changes.

    OSS layer: re-uses `knowledge.models.File` as the durable blob store.
    File.save(bytea) creates a Postgres large object, `get_bytes()` reads it
    back. The "OSS key" we persist on DocumentTemplate / DocumentGeneration
    is the File primary-key UUID (as string).
"""
from __future__ import annotations

import uuid as _uuid
from io import BytesIO
from typing import Optional
from uuid import UUID

import uuid_utils.compat as uuid

from common.utils.logger import maxkb_logger

from .perf import log_slow


_FILE_SOURCE_TYPE = 'SYSTEM'  # keeps finance docs from being garbage-collected by TEMPORARY_* sweeps
_FILE_SOURCE_ID_TEMPLATE = 'FINANCE_TEMPLATE'
_FILE_SOURCE_ID_GENERATION = 'FINANCE_GENERATION'


# --------------------------------------------------------------------------
# OSS helpers (thin wrappers — File model is the source of truth for storage).
# --------------------------------------------------------------------------


def _save_bytes(file_bytes: bytes, file_name: str, source_id: str) -> str:
    """
    Store `file_bytes` as a knowledge.models.File and return its id (as the
    OSS key we persist on the template / generation row).
    """
    from knowledge.models import File

    file_id = uuid.uuid7()
    f = File(
        id=file_id,
        file_name=file_name,
        meta={'finance': True, 'source_id': source_id},
        source_id=source_id,
        source_type=_FILE_SOURCE_TYPE,
    )
    f.save(file_bytes)
    return str(file_id)


def _load_bytes(oss_key: str) -> bytes:
    """Fetch bytes by OSS key. Raises ValueError on a missing key."""
    from knowledge.models import File

    try:
        # File.id is a UUID — coerce explicitly so a stray string format works.
        file_id = _uuid.UUID(str(oss_key))
    except (ValueError, AttributeError, TypeError) as e:
        raise ValueError(f'invalid oss key: {oss_key!r}') from e
    f = File.objects.filter(id=file_id).first()
    if f is None:
        raise ValueError(f'oss key not found: {oss_key}')
    return f.get_bytes()


# --------------------------------------------------------------------------
# Public service surface
# --------------------------------------------------------------------------


def store_template_bytes(file_bytes: bytes, file_name: str) -> str:
    """Persist a freshly uploaded template docx. Returns its OSS key."""
    return _save_bytes(file_bytes, file_name, _FILE_SOURCE_ID_TEMPLATE)


def load_template_bytes(oss_key: str) -> bytes:
    """Public re-export of `_load_bytes` for callers fetching a stored template."""
    return _load_bytes(oss_key)


def build_sample_template_bytes() -> bytes:
    """
    Generate a sample .docx demonstrating docxtpl Jinja-style placeholder syntax.

    Naming suffix → type-inference hints (see service/template_parser._infer_type):
      *_amount, *_count, *_number, *_qty, *_total       → number
      *_date, *_at, *_time, *_deadline, *_expiry        → date
      *_desc, *_description, *_summary, *_analysis,
      *_reason, *_notes, *_remark, *_content            → long_text
      anything else                                     → text

    The sample is built fresh on each call (no fixture file on disk) so its
    contents always reflect the current placeholder/type contract. Pure
    python-docx — no template engine involvement.
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    title = doc.add_heading('融资项目模板示例 / Finance Project Template — Sample', level=0)
    title.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro_run = intro.add_run(
        '本文件演示 docxtpl Jinja 风格占位符语法。把"双花括号包裹的变量名"'
        '（见下文正文示例）替换成你自己的 key，其余文字按需修改后保存即可上传。'
    )
    # 注意：**不能**在 prose 里写裸 `{{ ... }}` 之类的语法演示 ——
    # docxtpl 会把整份文档喂给 Jinja 解析，`...` 不是合法标识符，整个上传流程
    # 会被 `unexpected '.'` 卡住。所有 `{{ var }}` 字面量都必须是**真实**占位符。
    intro_run.font.size = Pt(10)

    # 设计说明（**勿删，给维护者**）：
    #
    # docxtpl 抽取变量时把整个 docx XML 喂给 Jinja2 env.parse(xml)，对 *任何*
    # 出现在文档里的 Jinja 语法都尝试解析 —— 包括"演示语法"的字面文本。一旦
    # python-docx 把 `{% %}` 块跨 <w:r> run 拆开（中文/英文/标点常触发），
    # 块尾尖括号或 dotted-access 会在 XML 标签夹缝中产出无效片段，
    # `unexpected '.'` / `unexpected '<'` 直接让上传失败。
    #
    # 经验上 **样本 docx 里只放 `{{ var }}` 单变量替换** 才稳定。`{% if %}` /
    # `{% for %}` / 过滤器 / dotted-access 这些进阶语法放到前端帮助面板讲
    # （那里只是 HTML 字符，不经 docxtpl 解析）。
    #
    # ---- 占位符使用说明（**纯文字**，不含任何 Jinja 字面量） ----
    doc.add_heading('一、占位符使用说明', level=1)
    p = doc.add_paragraph()
    p.add_run('1. 单变量替换：').bold = True
    p.add_run('用双花括号包裹变量名，前后各加一个空格。具体写法参见下方"三、模板正文"。')
    p2 = doc.add_paragraph()
    p2.add_run('2. 变量命名规则：').bold = True
    p2.add_run(
        '只允许字母、数字、下划线，首字符必须是字母。后缀决定推断的字段类型（详见第二节）。'
    )
    p3 = doc.add_paragraph()
    p3.add_run('3. 进阶语法（条件 / 循环 / 过滤器）：').bold = True
    p3.add_run(
        '本样本只演示单变量替换 —— 进阶语法跨 run 拆分时容易让 docxtpl 解析失败。'
        '需要时请参考上传页的"占位符语法帮助"面板或 docxtpl 官方文档。'
    )

    # ---- 类型推断表 ----
    doc.add_heading('二、命名后缀决定字段类型', level=1)
    typetable = doc.add_table(rows=1, cols=3)
    typetable.style = 'Light Grid Accent 1'
    th = typetable.rows[0].cells
    th[0].text = '后缀'
    th[1].text = '推断类型'
    th[2].text = '示例 key'
    type_rows = [
        ('_amount / _count / _number / _qty / _total', 'number', 'target_amount'),
        ('_date / _at / _time / _deadline / _expiry', 'date', 'report_date'),
        ('_desc / _description / _summary / _analysis / _reason / _notes / _remark / _content', 'long_text', 'risk_analysis'),
        ('（其他）', 'text', 'project_name'),
    ]
    for suf, typ, exa in type_rows:
        r = typetable.add_row().cells
        r[0].text = suf
        r[1].text = typ
        r[2].text = exa

    # ---- 可直接编辑的模板正文 ----
    doc.add_heading('三、模板正文（可直接修改后上传）', level=1)

    doc.add_heading('项目概览', level=2)
    doc.add_paragraph('项目名称：{{ project_name }}')
    doc.add_paragraph('借款主体：{{ borrower_name }}')
    doc.add_paragraph('融资类型：{{ project_type }}')
    doc.add_paragraph('行业代码：{{ industry_code }}')

    doc.add_heading('金额与时间', level=2)
    doc.add_paragraph('目标金额：{{ target_amount }} 元')
    doc.add_paragraph('期限月数：{{ duration_months }}')
    doc.add_paragraph('报告日期：{{ report_date }}')
    doc.add_paragraph('预计到账：{{ expected_close_date }}')

    doc.add_heading('描述与分析（长文本）', level=2)
    doc.add_paragraph('项目描述：{{ project_description }}')
    doc.add_paragraph('风险分析：{{ risk_analysis }}')
    doc.add_paragraph('还款来源说明：{{ repayment_source_notes }}')

    doc.add_heading('其他说明', level=2)
    doc.add_paragraph('抵押情况说明：{{ collateral_summary }}')
    doc.add_paragraph('材料清单概要：{{ materials_summary }}')

    foot = doc.add_paragraph()
    foot.add_run(
        '提示：保存为 .docx（不要存成 .doc）后回到上传页继续。'
        '占位符的标签、类型、是否必填可在上传后的详情页继续编辑。'
        '若要使用条件块、循环或属性访问等高级 Jinja 语法，请在 Word 里把整段表达式选中'
        '设为同一字体（同一 <w:r> run）后再上传，避免 docxtpl 跨 run 解析失败。'
    ).italic = True

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


@log_slow(threshold_ms=1000, name='finance.document_generator.render_template')
def render_template(template_oss_key: str, values: dict, output_filename: str) -> str:
    """
    Render the template at `template_oss_key` against `values` and store the
    resulting docx as a new File. Returns the new OSS key.
    """
    from docxtpl import DocxTemplate

    template_bytes = _load_bytes(template_oss_key)
    doc = DocxTemplate(BytesIO(template_bytes))
    # docxtpl tolerates unknown placeholders → renders them as empty.
    doc.render(values or {})

    out = BytesIO()
    doc.save(out)
    return _save_bytes(out.getvalue(), output_filename, _FILE_SOURCE_ID_GENERATION)


@log_slow(threshold_ms=1000, name='finance.document_generator.render_to_html')
def render_to_html(output_oss_key: str) -> str:
    """
    Convert a rendered docx (already in OSS) to an HTML fragment for preview.
    Returns just the body HTML — caller is expected to wrap it in their own
    page/iframe.
    """
    import mammoth

    docx_bytes = _load_bytes(output_oss_key)
    with BytesIO(docx_bytes) as buf:
        result = mammoth.convert_to_html(buf)
    return result.value or ''


@log_slow(threshold_ms=1000, name='finance.document_generator.trigger_generation')
def trigger_generation(generation_id: UUID) -> None:
    """
    Process one DocumentGeneration row through the state machine.

    Background-friendly: takes only an id, looks the row up itself, never
    raises. On success: status=PENDING_REVIEW, output_oss_key=<key>.
    On failure: status=FAILED, error_message=<repr(exception)>.

    Idempotency: caller is expected to invoke this exactly once per row
    while the row is in GENERATING. The function refuses to re-process rows
    in any other state.
    """
    from finance.models import DocumentGeneration, DocumentTemplate, GenerationStatus

    gen: Optional[DocumentGeneration] = DocumentGeneration.objects.filter(
        id=generation_id
    ).first()
    if gen is None:
        maxkb_logger.error(f'[finance.gen] generation {generation_id} not found')
        return
    if gen.status != GenerationStatus.GENERATING:
        maxkb_logger.warning(
            f'[finance.gen] skip {generation_id}: status={gen.status} (expected generating)'
        )
        return

    template = DocumentTemplate.objects.filter(id=gen.template_id, is_deleted=False).first()
    if template is None:
        gen.status = GenerationStatus.FAILED
        gen.error_message = f'template {gen.template_id} not found or deleted'
        gen.save(update_fields=['status', 'error_message', 'updated_at'])
        return

    try:
        output_filename = f'finance-{gen.id}.docx'
        output_key = render_template(
            template_oss_key=template.docx_oss_key,
            values=gen.placeholder_values or {},
            output_filename=output_filename,
        )
        gen.output_oss_key = output_key
        gen.status = GenerationStatus.PENDING_REVIEW
        gen.error_message = ''
        gen.save(
            update_fields=['output_oss_key', 'status', 'error_message', 'updated_at']
        )
    except Exception as e:  # noqa: BLE001 — surface anything as a FAILED row
        maxkb_logger.error(
            f'[finance.gen] render failed for {generation_id}: {e}', exc_info=True
        )
        gen.status = GenerationStatus.FAILED
        gen.error_message = repr(e)[:8000]
        gen.save(update_fields=['status', 'error_message', 'updated_at'])
